import os
from langchain_core.tools import tool
from pypdf import PdfReader
from docx import Document

def resolve_file(file_path: str) -> str:
    """Mencari lokasi file yang valid di direktori lokal atau temp_uploads."""
    if os.path.exists(file_path):
        return file_path
    
    # Cek di folder temp_uploads
    temp_dir = os.path.join(os.getcwd(), "temp_uploads")
    if os.path.exists(temp_dir):
        # Cek exact basename
        candidate = os.path.join(temp_dir, os.path.basename(file_path))
        if os.path.exists(candidate):
            return candidate
        # Cek kecocokan nama file
        for f in os.listdir(temp_dir):
            if os.path.basename(file_path).lower() in f.lower() or f.lower() in os.path.basename(file_path).lower():
                return os.path.join(temp_dir, f)
    return file_path

@tool
def read_pdf(file_path: str, max_pages: int = 20) -> str:
    """Membaca dan mengekstrak teks dari file dokumen PDF.
    Gunakan tool ini ketika perlu membaca soal latihan, petunjuk, atau teori dari file PDF.
    """
    real_path = resolve_file(file_path)
    if not os.path.exists(real_path):
        return f"Error: File PDF '{file_path}' tidak ditemukan di sistem."
    
    try:
        reader = PdfReader(real_path)
        total_pages = len(reader.pages)
        pages_to_read = min(total_pages, max_pages)
        
        extracted_text = []
        for i in range(pages_to_read):
            page_text = reader.pages[i].extract_text()
            if page_text:
                extracted_text.append(f"--- [Halaman {i+1} dari {total_pages}] ---\n{page_text.strip()}")
        
        if not extracted_text:
            return f"Peringatan: File PDF '{os.path.basename(real_path)}' tidak berisi teks yang dapat diekstrak (kemungkinan berupa scan/gambar murni)."
        
        result = "\n\n".join(extracted_text)
        if total_pages > max_pages:
            result += f"\n\n[Catatan: Hanya {max_pages} halaman pertama yang diekstrak dari total {total_pages} halaman.]"
        return result
    except Exception as e:
        return f"Error saat membaca file PDF: {str(e)}"

@tool
def read_word(file_path: str) -> str:
    """Membaca dan mengekstrak seluruh isi teks dan tabel dari file Microsoft Word (.docx).
    Gunakan tool ini ketika perlu membaca soal atau deskripsi tugas dari file Word.
    """
    real_path = resolve_file(file_path)
    if not os.path.exists(real_path):
        return f"Error: File Word '{file_path}' tidak ditemukan di sistem."
    
    try:
        doc = Document(real_path)
        extracted_content = []
        
        # Ekstrak paragraf
        for p in doc.paragraphs:
            if p.text.strip():
                extracted_content.append(p.text.strip())
        
        # Ekstrak tabel jika ada
        for t_idx, table in enumerate(doc.tables):
            extracted_content.append(f"\n[Tabel {t_idx + 1}]")
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                extracted_content.append(" | ".join(row_data))
        
        if not extracted_content:
            return f"Peringatan: Dokumen Word '{os.path.basename(real_path)}' kosong."
            
        return "\n".join(extracted_content)
    except Exception as e:
        return f"Error saat membaca file Word: {str(e)}"
